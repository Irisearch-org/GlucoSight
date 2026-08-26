from __future__ import annotations

import uuid
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd

SOURCE_DATASET = "shanghai_t2dm"
SCHEMA_VERSION = "1.2"

DATE_CANDIDATES = {"Date"}
CGM_CANDIDATES = {"CGM (mg / dl)"}
DIET_EN_CANDIDATES = {"Dietary intake"}
DIET_CN_CANDIDATES = {"饮食", "进食量"}
CBG_CANDIDATES = {"CBG (mg / dl)"}

_MISSING_TOKENS_LOWER = {
    "", "/", "-", "—", "–",
    "not available", "data not available", "data  not available",
    "na", "n/a", "nan", "none", "null",
    "未记录", "未进食", "无",
}

_HERE = Path(__file__).resolve().parent
DEFAULT_DATA_DIR = _HERE / "Shanghai_T2DM"


def _is_missing(value) -> bool:
    if value is None:
        return True
    s = str(value).strip()
    if s in ("", "/", "-", "—", "–", "未记录", "未进食", "无"):
        return True
    if s.lower() in _MISSING_TOKENS_LOWER:
        return True
    if " ".join(s.lower().split()) in _MISSING_TOKENS_LOWER:
        return True
    return s in _MISSING_TOKENS_LOWER


def _find_col(header: List[str], candidates: set) -> Optional[int]:
    for i, h in enumerate(header):
        if h is None:
            continue
        hs = str(h).strip()
        if hs in candidates or hs.lower() in {c.lower() for c in candidates}:
            return i
    return None


def _parse_timestamp(value, datemode: Optional[int] = None) -> Optional[datetime]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, (int, float)) and datemode is not None:
        try:
            import xlrd
            return xlrd.xldate_as_datetime(float(value), datemode)
        except Exception:
            return None
    s = str(value).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        return None


def _meal_id(participant_id: str, t0: datetime, diet_text: str) -> str:
    base = f"{participant_id}|{t0.isoformat()}|{diet_text.strip()}"
    return str(uuid.uuid5(uuid.NAMESPACE_DNS, base))


def _read_excel(path: Path) -> Tuple[List[str], List[Tuple], Optional[int]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], [], None
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        return header, rows[1:], None
    if suffix == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(path))
        sh = wb.sheet_by_index(0)
        header = [str(sh.cell_value(0, c)).strip() for c in range(sh.ncols)]
        data = [tuple(sh.cell_value(r, c) for c in range(sh.ncols)) for r in range(1, sh.nrows)]
        return header, data, wb.datemode
    raise ValueError(f"Unsupported suffix: {path}")


def load_shanghai_diet_records(data_dir: Optional[Path] = None, *, generate_meal_ids: bool = True) -> pd.DataFrame:
    if data_dir is None:
        data_dir = DEFAULT_DATA_DIR
    data_dir = Path(data_dir)
    if not data_dir.exists():
        raise FileNotFoundError(f"ShanghaiT2DM dir not found: {data_dir}")

    files = sorted(set(sorted(data_dir.glob("*.xls")) + sorted(data_dir.glob("*.xlsx"))))
    if not files:
        raise FileNotFoundError(f"No Excel files in {data_dir}")

    records: List[Dict] = []
    for path in files:
        parts = path.stem.split("_")
        participant_id = parts[0] if parts else path.stem
        visit_id = parts[1] if len(parts) > 1 else ""
        header, data_rows, datemode = _read_excel(path)
        if not header:
            continue

        idx_date = _find_col(header, DATE_CANDIDATES) or 0
        idx_cgm = _find_col(header, CGM_CANDIDATES)
        idx_cbg = _find_col(header, CBG_CANDIDATES)
        idx_en = _find_col(header, DIET_EN_CANDIDATES)
        if idx_en is None:
            idx_en = 4 if len(header) > 4 else None
        idx_cn = _find_col(header, DIET_CN_CANDIDATES)

        for r_idx, row in enumerate(data_rows, start=2):
            diet_en_raw = row[idx_en] if idx_en is not None and idx_en < len(row) else None
            if _is_missing(diet_en_raw):
                continue
            diet_en = str(diet_en_raw).strip()
            diet_cn_raw = row[idx_cn] if idx_cn is not None and idx_cn < len(row) else None
            diet_cn = "" if _is_missing(diet_cn_raw) else str(diet_cn_raw).strip()

            raw_date = row[idx_date] if idx_date < len(row) else None
            t0 = _parse_timestamp(raw_date, datemode)
            if t0 is None:
                continue

            cgm_val = None
            if idx_cgm is not None and idx_cgm < len(row):
                v = row[idx_cgm]
                if v is not None and str(v).strip() not in ("", "/", "-", "—"):
                    try:
                        cgm_val = float(v)
                    except (TypeError, ValueError):
                        pass

            cbg_val = None
            if idx_cbg is not None and idx_cbg < len(row):
                v = row[idx_cbg]
                if v is not None and str(v).strip() not in ("", "/", "-", "—"):
                    try:
                        cbg_val = float(v)
                    except (TypeError, ValueError):
                        pass

            records.append({
                "meal_id": _meal_id(participant_id, t0, diet_en) if generate_meal_ids else "",
                "participant_id": str(participant_id),
                "visit_id": str(visit_id),
                "file_name": path.name,
                "row_index": r_idx,
                "t0_timestamp": t0,
                "t0_timestamp_iso": t0.isoformat(),
                "source_dataset": SOURCE_DATASET,
                "schema_version": SCHEMA_VERSION,
                "dietary_intake": diet_en,
                "dietary_intake_cn": diet_cn,
                "cgm_mgdl": cgm_val,
                "cbg_mgdl": cbg_val,
                "cgm_available": 1 if cgm_val is not None else 0,
            })

    df = pd.DataFrame(records)
    if not df.empty:
        df.sort_values(by=["participant_id", "t0_timestamp"], inplace=True, kind="mergesort")
        df.reset_index(drop=True, inplace=True)
    return df


def inspect_shanghai_dataset(df: Optional[pd.DataFrame] = None, data_dir: Optional[Path] = None) -> Dict:
    if df is None:
        df = load_shanghai_diet_records(data_dir=data_dir)
    if df.empty:
        print("[inspect] No records found.")
        return {"n_records": 0, "has_variation": False}

    n_records = len(df)
    n_participants = df["participant_id"].nunique()
    n_files = df["file_name"].nunique()
    n_unique_en = df["dietary_intake"].nunique()
    n_unique_cn = df["dietary_intake_cn"].nunique()

    counter = Counter(df["dietary_intake"])
    top_en = counter.most_common(5)
    top_share = top_en[0][1] / n_records if top_en else 0

    en_texts = df["dietary_intake"].tolist()
    cn_texts = df["dietary_intake_cn"].tolist()

    en_methods = {kw: sum(1 for t in en_texts if kw.lower() in t.lower()) for kw in ["fried", "boiled", "steamed", "grilled", "baked", "stir"]}
    cn_methods = {ch: sum(1 for t in cn_texts if ch in t) for ch in ["炒", "煮", "蒸", "煎", "炸", "烤", "拌", "炖", "烧"]}
    avg_len = sum(len(t) for t in en_texts) / len(en_texts) if en_texts else 0
    has_variation = n_unique_en > 500 and avg_len > 20 and top_share < 0.10 and (sum(en_methods.values()) + sum(cn_methods.values()) > 100)

    print("=" * 70)
    print("SHANGHAIT2DM DIETARY RECORD INSPECTION")
    print("=" * 70)
    print(f"Files scanned          : {n_files}")
    print(f"Participants           : {n_participants}")
    print(f"Non-empty records      : {n_records}")
    print(f"Unique EN / CN         : {n_unique_en} / {n_unique_cn}")
    print(f"Avg EN length          : {avg_len:.1f}")
    print(f"Top EN share           : {top_share:.1%} ({top_en[0][0][:60]!r})" if top_en else "Top EN share: —")
    print(f"Top 5 EN:")
    for txt, cnt in top_en:
        print(f"  {cnt:4d}x  {txt.replace(chr(10), ' / ')[:80]!r}")
    print(f"EN methods             : {en_methods}")
    print(f"CN methods             : {cn_methods}")
    print("-" * 70)
    if has_variation:
        print("GATE: PASS — real variation, Path A is genuine extraction task.")
    else:
        print("GATE: COLLAPSE — uniform/thin, Path A -> lookup, focus Path B.")
    print("=" * 70)
    print(f"Contract: {SOURCE_DATASET} / {SCHEMA_VERSION}, missing excluded (H1).")
    return {
        "n_files": n_files, "n_participants": n_participants, "n_records": n_records,
        "n_unique_en": n_unique_en, "n_unique_cn": n_unique_cn,
        "top_en": top_en, "avg_len": avg_len, "has_variation": has_variation,
        "en_methods": en_methods, "cn_methods": cn_methods,
    }


if __name__ == "__main__":
    import argparse
    try:
        import sys
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass
    parser = argparse.ArgumentParser(description="ShanghaiT2DM dietary loader — v1.2")
    parser.add_argument("--data-dir", type=str, default=None)
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--head", type=int, default=5)
    parser.add_argument("--save", type=str, default=None)
    args = parser.parse_args()

    data_dir = Path(args.data_dir) if args.data_dir else DEFAULT_DATA_DIR
    print(f"Scanning: {data_dir}")
    df = load_shanghai_diet_records(data_dir=data_dir)
    if df.empty:
        print("No records.")
    else:
        print(f"Loaded {len(df)} records from {df['file_name'].nunique()} files, {df['participant_id'].nunique()} participants.")
        print(df.head(args.head).to_string(index=False))
        if args.inspect:
            inspect_shanghai_dataset(df)
        if args.save:
            out = Path(args.save)
            out.parent.mkdir(parents=True, exist_ok=True)
            df.to_csv(out, index=False, encoding="utf-8-sig")
            print(f"Saved to {out}")
